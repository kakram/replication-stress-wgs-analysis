#!/usr/bin/env python3
"""
Variant burden table generator — MCF-7 somatic variant burden as a Word
table (Table 5.X; formerly 5.3.1).

Reads outputs/ch5_variant_burden.csv (primary-contig counts by default)
and writes two .docx files, each containing a native Word table that can
be copied into the main thesis document:

  outputs/Table_5_X_variant_burden_full.docx       (exhaustive reference)
  outputs/Table_5_X_variant_burden_condensed.docx  (main-text version)

Changes 2026-09-09: TMB row removed (provider TMB has an undocumented
definition and is not reported); "Excluded: non-primary contigs" row
added; caption states the calling design and the contig policy.

Usage:
    python3 scripts/make_ch5_variant_burden_table.py [--csv PATH]

Requires:
    pip install python-docx pandas
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    sys.exit(
        "python-docx not installed. Install with:\n"
        "    pip install python-docx"
    )


REPO = Path(__file__).resolve().parent.parent
CSV_IN = REPO / "outputs" / "ch5_variant_burden.csv"
DOCX_FULL = REPO / "outputs" / "Table_5_X_variant_burden_full.docx"
DOCX_CONDENSED = REPO / "outputs" / "Table_5_X_variant_burden_condensed.docx"

SAMPLE_COLS = ["WTUN", "WTAPH", "B4UN", "B4APH"]
SAMPLE_HEADERS = {
    "WTUN":  "WT untreated",
    "WTAPH": "WT + aphidicolin",
    "B4UN":  "B4 ZFP36L1\u207B\u002F\u207B untreated",
    "B4APH": "B4 ZFP36L1\u207B\u002F\u207B + aphidicolin",
}

CAPTION = (
    "Table 5.X. Genome-wide somatic small-variant burden across four MCF-7 "
    "conditions. Wild-type samples were called in tumour-only mode and B4 "
    "samples in paired tumour\u2013normal mode against the treatment-matched "
    "wild-type, so counts are comparable within, but not between, "
    "genotypes. Counts are restricted to the primary assembly (chr1\u201322, "
    "X, Y, M); records on unplaced, alternate, decoy and HLA contigs are "
    "excluded and their number is shown. Filter labels follow GATK/Mutect2 "
    "conventions; rows carrying multiple filter terms increment each "
    "matching filter column once, so column sums may exceed "
    "(Total \u2212 PASS)."
)

# Each entry is (display_label, csv_column, is_subrow_indented)
FULL_ROWS = [
    ("Total variants (primary contigs)",     "total_variants",            False),
    ("Excluded: non-primary contigs",        "vcf_records_nonprimary",    True),
    ("PASS variants",                        "pass_variants",             False),
    ("FILTER: germline",                     "filter_germline",           True),
    ("FILTER: panel_of_normals",             "filter_panel_of_normals",   True),
    ("FILTER: weak_evidence",                "filter_weak_evidence",      True),
    ("FILTER: clustered_events",             "filter_clustered",          True),
    ("FILTER: other",                        "filter_other",              True),
    ("Coding variants (total)",              "coding_total",              False),
    ("Missense",                             "missense",                  True),
    ("Synonymous",                           "synonymous",                True),
    ("Nonsense",                             "nonsense",                  True),
    ("Frameshift deletion",                  "frameshift_del",            True),
    ("Frameshift insertion",                 "frameshift_ins",            True),
    ("In-frame deletion",                    "in_frame_del",              True),
    ("In-frame insertion",                   "in_frame_ins",              True),
    ("Splice site",                          "splice_site",               True),
    ("Splice region",                        "splice_region",             True),
    ("SIFT deleterious",                     "sift_deleterious",          False),
    ("SIFT tolerated",                       "sift_tolerated",            False),
    ("SIFT unknown",                         "sift_unknown",              False),
    ("gnomAD known (AF > 0.001)",            "gnomad_known",              False),
    ("gnomAD novel (AF \u2264 0.001)",        "gnomad_novel",              False),
    ("Genes with \u22651 variant",           "genes_with_any_variant",    False),
    ("Genes with coding variant",            "genes_with_coding_variant", False),
]

# Condensed version — the rows referenced in the §5.3 prose.
CONDENSED_KEEP = {
    "total_variants",
    "vcf_records_nonprimary",
    "pass_variants",
    "filter_germline",
    "filter_panel_of_normals",
    "filter_weak_evidence",
    "filter_clustered",
    "filter_other",
    "coding_total",
    "missense",
    "sift_deleterious",
    "sift_tolerated",
    "gnomad_novel",
}


def fmt(value, col: str) -> str:
    """Thousands-separated integer; floats (none expected now) to 2 d.p."""
    if isinstance(value, float) and not float(value).is_integer():
        return f"{value:.2f}"
    return f"{int(value):,}"


def build_table(doc, rows, df):
    n_cols = 1 + len(SAMPLE_COLS)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"
    # Fixed layout: 5.5 cm metric column, remainder shared by the samples
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    widths = [Cm(5.5)] + [Cm((16.0 - 5.5) / len(SAMPLE_COLS))] * len(SAMPLE_COLS)
    for ci, col in enumerate(table.columns):
        col.width = widths[ci]
    for r in table.rows:
        for ci, c in enumerate(r.cells):
            c.width = widths[ci]

    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    for i, s in enumerate(SAMPLE_COLS):
        hdr[i + 1].text = SAMPLE_HEADERS[s]
    for cell in hdr:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    for i, (label, col, indent) in enumerate(rows, start=1):
        row = table.rows[i].cells
        prefix = "    " if indent else ""
        row[0].text = prefix + label
        for j, s in enumerate(SAMPLE_COLS):
            value = df.loc[df["sample"] == s, col].iloc[0]
            row[j + 1].text = fmt(value, col)
            for p in row[j + 1].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for r in table.rows:
        for cell in r.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)


def write_docx(rows, out_path: Path, df: pd.DataFrame):
    doc = Document()
    cap = doc.add_paragraph()
    cap_run = cap.add_run(CAPTION)
    cap_run.font.size = Pt(10)
    doc.add_paragraph()
    build_table(doc, rows, df)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"  wrote {out_path.relative_to(REPO)}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", type=Path, default=CSV_IN)
    args = ap.parse_args()

    if not args.csv.exists():
        sys.exit(f"Input CSV not found: {args.csv}")

    df = pd.read_csv(args.csv)
    if "vcf_records_nonprimary" not in df.columns:
        sys.exit("CSV predates the primary-contig restriction; re-run "
                 "scripts/ch5_variant_burden.py first.")
    if (df["contig_set"] != "primary").any():
        print("WARNING: CSV was generated with --all-contigs; the caption "
              "states primary contigs. Regenerate or edit the caption.")

    print("=== Building variant burden table (full and condensed) ===\n")
    print(f"Loaded {len(df)} samples from {args.csv.relative_to(REPO)}\n")

    write_docx(FULL_ROWS, DOCX_FULL, df)
    condensed = [r for r in FULL_ROWS if r[1] in CONDENSED_KEEP]
    write_docx(condensed, DOCX_CONDENSED, df)

    print("\nDone. Open either .docx in Word, copy the table, paste into the thesis.")


if __name__ == "__main__":
    main()
